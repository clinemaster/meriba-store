"""Historical Excel import and review: every exception the business owner listed on 12 Oct 2026,
plus dry run, traceability, idempotency and separation from the live system (AT-16, AT-17, AT-STK-06)."""
import os
import uuid
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from catalog.models import Product
from core.models import AuditLog, Branch
from history.models import (ApprovedPrice, HistoricalRecord, HistoryAcceptance, ImportBatch, ReviewItem, ReviewNote, SourceProduct,
                            UnclassifiedRow)
from inventory.models import InventoryMovement, StockBalance
from sales.models import Payment, Sale
from trading.models import BusinessDay

from .conftest import client_for, sell, start_shift, stock_of

HEADER = ["S/N", "Details", "2", "Add Purchases", "Total  stock", "Sales", "Closing stock", "Purchase  price", "Sales  price",
          "Total  Sales", "Profit", "Stock  ", "T.Purchases"]
AUG31, SEP10, SEP11, SEP12 = date(2026, 8, 31), date(2026, 9, 10), date(2026, 9, 11), date(2026, 9, 12)


def line(serial, name, opening, sold, cost, price, bought=None):
    closing = opening + (bought or 0) - sold
    return [serial, name, opening, bought, opening + (bought or 0), sold or None, closing, cost, price, sold * price,
            sold * (price - cost), closing * price, (bought or 0) * cost]


_BUILT: dict[str, bytes] = {}


def workbook(days: dict, extra_sheets=(), with_totals=True) -> bytes:
    """Builds a workbook laid out like the shop's: one sheet per day, a title, a heading row, products, then unlabelled totals."""
    # A saved workbook carries the time it was saved, so the same content built a second later is a different file.
    # The same content must give the same file here, as it would when the shop uploads one file twice.
    key = repr((days, extra_sheets, with_totals))
    if key in _BUILT:
        return _BUILT[key]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for day, rows in days.items():
        ws = wb.create_sheet(day.strftime("%d.%m.%Y"))
        ws.append([])
        ws.append([None, "MERIBA   STORE"])
        ws.append(HEADER)
        for row in rows:
            ws.append(row)
        if with_totals:
            ws.append([])
            totals = [None] * 13
            totals[9], totals[10] = sum(r[9] for r in rows), sum(r[10] for r in rows)  # sums of Total Sales and Profit
            ws.append(totals)
            ws.append([None] * 9 + [632700])  # a figure with no label and no match
    for name in extra_sheets:
        wb.create_sheet(name).append(["notes", "not a daily sheet"])
    out = BytesIO()
    wb.save(out)
    _BUILT[key] = out.getvalue()
    return _BUILT[key]


def primary() -> bytes:
    """31 Aug, then 10-12 Sep (1-9 Sep missing). Carries every exception the owner asked to be handled."""
    def day(konyagi_open, konyagi_sold, magic_cost):
        return [
            line(1, "Konyagi  ndogo", konyagi_open, konyagi_sold, 3750, 4500),
            line(2, "Henken", 102, 0, 2917, 3500),
            line(3, "Henken", 0, 0, 3125, 4000),
            line(4, "MAGIC MOMENT kati", 48, 0, magic_cost, 13000),
            line(5, "mkwawa", 0, 0, 0, 0),
            line(6, "Safari Kubwa", 20, 0, 34000, 36000),
        ]
    return workbook({AUG31: day(100, 10, 9000), SEP10: day(70, 5, 9000), SEP11: day(65, 5, 9000), SEP12: day(60, 0, 1290)})


def upload(client, data: bytes, url="/api/v1/history/imports/", name="MERIBA STORE.xlsx", **fields):
    file = SimpleUploadedFile(name, data, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    return client.post(url, {"file": file, **fields}, format="multipart")


def commit(client, data: bytes, **kw):
    return upload(client, data, confirm="true", backup_confirmed="true", **kw)


@pytest.fixture
def imported(owner):
    response = commit(client_for(owner), primary())
    assert response.status_code == 201, response.data
    return response.data


def item(kind) -> ReviewItem:
    return ReviewItem.objects.get(kind=kind)


def decide(client, review_item, resolution, note="Checked against the stock book with the owner", **extra):
    return client.post(f"/api/v1/history/review-items/{review_item.pk}/decide/", {"resolution": resolution, "note": note} | extra, format="json")


# --- Dry run, commit, traceability, idempotency ----------------------------------------------

def test_dry_run_reports_everything_and_saves_nothing(owner, accountant, cashier):
    assert upload(client_for(cashier), primary(), "/api/v1/history/imports/dry-run/").status_code == 403
    response = upload(client_for(accountant), primary(), "/api/v1/history/imports/dry-run/")
    assert response.status_code == 200, response.data
    summary = response.data["summary"]
    assert (summary["worksheets_read"], summary["product_rows"], summary["accepted_new"], summary["duplicates_skipped"]) == (4, 24, 24, 0)
    assert (summary["date_from"], summary["date_to"], summary["unlabelled_rows"], summary["source_products"]) == ("2026-08-31", "2026-09-12", 8, 6)
    assert summary["workbook_total_sales"] == "90000"  # 20 Konyagi at 4,500: checked against the workbook itself (FR-IMP-12)
    kinds = sorted(e["kind"] for e in response.data["exceptions"])
    assert kinds == ["MISSING_DATES", "MISSING_PRICE", "PRICE_ANOMALY", "SAME_NAME_DIFFERENT_PRICE"]
    assert all(e["is_new"] for e in response.data["exceptions"])
    assert "Nothing has been saved" in response.data["statement"]
    for model in (ImportBatch, HistoricalRecord, SourceProduct, ReviewItem, UnclassifiedRow):
        assert not model.objects.exists(), model.__name__


def test_commit_is_owner_only_and_needs_backup_and_preview_confirmation(owner, accountant):
    assert commit(client_for(accountant), primary()).status_code == 403
    boss = client_for(owner)
    assert upload(boss, primary(), confirm="true").data["code"] == "backup_required"
    assert upload(boss, primary(), backup_confirmed="true").data["code"] == "not_confirmed"
    assert upload(boss, b"not a workbook", confirm="true", backup_confirmed="true").data["code"] == "unreadable"
    assert upload(boss, primary(), name="old.xls", confirm="true", backup_confirmed="true").data["code"] == "bad_file"
    assert not HistoricalRecord.objects.exists()


def test_records_keep_their_source_and_original_values(imported):
    assert imported["batch"]["summary"]["records_stored"] == 24 and imported["batch"]["summary"]["verified_against_workbook"] is True
    record = HistoricalRecord.objects.get(business_date=SEP10, source_product__serial_number=1)
    assert (record.source_file, record.worksheet, record.source_row) == ("MERIBA STORE.xlsx", "10.09.2026", 4)
    assert record.original["Details"] == "Konyagi ndogo" and record.original["Sales price"] == "4500" and record.original["Purchase price"] == "3750"
    assert (record.opening_qty, record.sold_qty, record.closing_qty, record.sales_price, record.purchase_price) == (70, 5, 65, 4500, 3750)
    assert record.validation_status == "VALID"
    empty = HistoricalRecord.objects.get(business_date=SEP10, source_product__serial_number=2)
    assert empty.sold_qty is None and empty.purchases_qty is None  # an empty cell stays empty; it is not turned into a zero


def test_at16_historical_import_never_touches_the_live_system(owner, cashier, shop, make_product):
    live = make_product("Konyagi ndogo", "4500", 48)
    till = client_for(cashier)
    start_shift(till)
    sell(till, [(live, 3)], [(shop.cash, 13500)])
    before = (stock_of(live), InventoryMovement.objects.count(), Sale.objects.count(), Payment.objects.count(),
              BusinessDay.objects.count(), StockBalance.objects.count(), Product.objects.count(), Branch.objects.count())
    assert commit(client_for(owner), primary()).status_code == 201
    after = (stock_of(live), InventoryMovement.objects.count(), Sale.objects.count(), Payment.objects.count(),
             BusinessDay.objects.count(), StockBalance.objects.count(), Product.objects.count(), Branch.objects.count())
    assert before == after  # AT-STK-06: historical closing figures are not added to live stock; no branch from a sheet name
    assert HistoricalRecord.objects.count() == 24


def test_at17_importing_the_same_workbook_again_adds_nothing(owner, imported):
    counts = lambda: (HistoricalRecord.objects.count(), SourceProduct.objects.count(), ReviewItem.objects.count(),
                      UnclassifiedRow.objects.count(), ImportBatch.objects.count())
    before = counts()
    again = commit(client_for(owner), primary())
    assert again.status_code == 200 and again.data["nothing_new"] is True
    assert again.data["preview"]["summary"]["duplicates_skipped"] == 24 and again.data["preview"]["summary"]["accepted_new"] == 0
    assert again.data["preview"]["summary"]["already_imported_file"] is True
    assert counts() == before


# --- 2. Missing dates and unexplained stock changes -------------------------------------------

def test_missing_dates_are_flagged_never_filled_in(owner, accountant, imported):
    gap = item("MISSING_DATES")
    assert gap.severity == "MATERIAL" and gap.status == "PENDING_REVIEW"
    assert gap.details["missing_dates"] == [str(date(2026, 9, 1) + timedelta(days=i)) for i in range(9)]
    affected = gap.details["affected_products"]
    assert [(a["product"], a["closing_before"], a["opening_after"], a["difference"]) for a in affected] == [("Konyagi ndogo", "90", "70", "-20")]
    assert affected[0]["before"] == {"worksheet": "31.08.2026", "row": 4, "date": "2026-08-31"}  # traceable to sheet and row
    assert "not assumed to be sales" in gap.details["statement"]
    # nothing was invented for the missing days
    assert not HistoricalRecord.objects.filter(business_date__gt=AUG31, business_date__lt=SEP10).exists()
    assert set(HistoricalRecord.objects.values_list("business_date", flat=True)) == {AUG31, SEP10, SEP11, SEP12}

    konyagi = HistoricalRecord.objects.get(business_date=SEP10, source_product__serial_number=1)
    safari = HistoricalRecord.objects.get(business_date=SEP10, source_product__serial_number=6)
    assert (konyagi.import_status, konyagi.incomplete) == ("PENDING_REVIEW", True)  # incomplete until reviewed
    assert (safari.import_status, safari.incomplete) == ("IMPORTED", False)  # its stock did not change across the gap

    office = client_for(accountant)
    note = office.post(f"/api/v1/history/review-items/{gap.pk}/notes/", {"text": "Depot delivery book shows 20 bottles returned on 4 Sep",
                       "evidence_reference": "Delivery book p.14", "source_reference": "Supplier credit note CN-221"}, format="json")
    assert note.status_code == 200 and note.data["notes"][0]["author"] == "Accountant"
    saved = ReviewNote.objects.get()
    assert (saved.evidence_reference, saved.source_reference) == ("Delivery book p.14", "Supplier credit note CN-221") and saved.created_at
    assert decide(office, gap, "EXPLAINED").data["code"] == "evidence_required"
    assert decide(office, gap, "EXPLAINED", note="short").data["code"] == "note_required"
    done = decide(office, gap, "EXPLAINED", evidence_reference="Delivery book p.14")
    assert done.status_code == 200 and done.data["status"] == "APPROVED" and done.data["decided_by_name"] == "Accountant"
    konyagi.refresh_from_db()
    assert (konyagi.import_status, konyagi.incomplete) == ("APPROVED", False)
    assert (konyagi.opening_qty, konyagi.original["Opening (column before Add Purchases)"]) == (70, "70")  # source values untouched
    assert {"history.note_added", "history.exception_decided"} <= set(AuditLog.objects.values_list("action", flat=True))


def test_accepting_a_gap_without_explanation_keeps_the_figures_marked_incomplete(owner, imported):
    gap = item("MISSING_DATES")
    assert decide(client_for(owner), gap, "ACCEPTED_INCOMPLETE").status_code == 200
    konyagi = HistoricalRecord.objects.get(business_date=SEP10, source_product__serial_number=1)
    assert (konyagi.import_status, konyagi.incomplete) == ("APPROVED", True)


# --- 3. Henken: same name, two prices ----------------------------------------------------------

def test_henken_lines_are_kept_apart_until_the_owner_decides(owner, accountant, imported):
    conflict = item("SAME_NAME_DIFFERENT_PRICE")
    lines = conflict.details["lines"]
    assert [(l["serial"], l["sales_price_history"][0]["amount"], l["purchase_price_history"][0]["amount"]) for l in lines] == [
        (2, "3500", "2917"), (3, "4000", "3125")]
    assert lines[0]["sales_price_history"][0] == {"amount": "3500", "from": "2026-08-31", "to": "2026-09-12"}  # source dates kept
    henkens = list(SourceProduct.objects.filter(normalised_name="henken").order_by("serial_number"))
    assert len(henkens) == 2 and all(h.product is None and h.review_status == "PENDING_REVIEW" for h in henkens)  # not merged
    assert HistoricalRecord.objects.filter(source_product__in=henkens).count() == 8  # both original records, every day
    assert {r["code"] for r in client_for(owner).get("/api/v1/history/review-items/?kind=SAME_NAME_DIFFERENT_PRICE").data["results"][0]["allowed_resolutions"]} == {
        "SAME_PRODUCT_PRICE_CHANGE", "DIFFERENT_VARIANTS", "DIFFERENT_PRICE_TYPES", "DATA_ENTRY_ERROR"}

    # the Accountant may comment but the decision is the Owner's
    refused = decide(client_for(accountant), conflict, "DIFFERENT_VARIANTS")
    assert refused.status_code == 403 and refused.data["code"] == "owner_only"
    boss = client_for(owner)
    blocked = boss.post(f"/api/v1/history/source-products/{henkens[0].pk}/decide/", {"approve": True, "confirmed_name": "Heineken", "confirmed_unit": "Bottle",
                        "create_product": True, "selling_price": "3500"}, format="json")
    assert blocked.status_code == 409 and blocked.data["code"] == "open_exceptions" and not Product.objects.exists()

    # a note alone cannot close it: identity, package size, price type and effective date are required for each line
    names_only = [{"source_product": h.pk, "confirmed_name": "Heineken"} for h in henkens]
    both = [{"source_product": h.pk, "confirmed_name": "Heineken", "confirmed_package_size": "330ml"} for h in henkens]
    one_price = [{"source_product": henkens[0].pk, "price_type": "SELLING", "amount": "3500", "effective_date": "2026-08-31"}]
    assert decide(boss, conflict, "DIFFERENT_VARIANTS").data["code"] == "identity_required"
    assert decide(boss, conflict, "DIFFERENT_VARIANTS", identities=names_only).data["code"] == "identity_required"
    assert decide(boss, conflict, "DIFFERENT_VARIANTS", identities=both).data["code"] == "price_required"
    assert decide(boss, conflict, "DIFFERENT_VARIANTS", identities=both, approved_prices=one_price).data["code"] == "price_required"
    no_date = decide(boss, conflict, "DIFFERENT_VARIANTS", identities=both, approved_prices=[
        {"source_product": h.pk, "price_type": "SELLING", "amount": "3500"} for h in henkens])
    assert no_date.status_code == 400 and "effective_date" in str(no_date.data["errors"])
    conflict.refresh_from_db()
    assert conflict.status == "PENDING_REVIEW" and not ApprovedPrice.objects.exists()

    done = decide(boss, conflict, "DIFFERENT_VARIANTS", note="Row 4 is the 330ml bottle, row 5 the 500ml can", identities=[
        {"source_product": henkens[0].pk, "confirmed_name": "Heineken bottle", "confirmed_package_size": "330ml"},
        {"source_product": henkens[1].pk, "confirmed_name": "Heineken can", "confirmed_package_size": "500ml"}], approved_prices=[
        {"source_product": henkens[0].pk, "price_type": "SELLING", "amount": "3500", "effective_date": "2026-08-31"},
        {"source_product": henkens[1].pk, "price_type": "SELLING", "amount": "4000", "effective_date": "2026-08-31"}])
    assert done.status_code == 200 and done.data["status"] == "APPROVED" and done.data["resolution_label"] == "Different package sizes or variants"
    prices = {(p.source_product.serial_number, p.price_type, p.amount, p.effective_date, p.approved_by) for p in ApprovedPrice.objects.all()}
    assert prices == {(2, "SELLING", Decimal("3500"), AUG31, owner), (3, "SELLING", Decimal("4000"), AUG31, owner)}
    henkens[0].refresh_from_db()
    assert (henkens[0].confirmed_name, henkens[0].confirmed_package_size) == ("Heineken bottle", "330ml")
    log = AuditLog.objects.get(action="history.exception_decided")
    assert log.actor == owner and log.after["resolution"] == "DIFFERENT_VARIANTS" and len(log.after["approved_prices"]) == 2
    # the originals are still exactly as imported
    assert set(HistoricalRecord.objects.filter(source_product__in=henkens).values_list("sales_price", flat=True)) == {Decimal("3500"), Decimal("4000")}


# --- 4. MAGIC MOMENT kati: unusual change in the purchase price --------------------------------

def test_price_anomaly_is_flagged_and_never_corrected_automatically(owner, imported):
    anomaly = item("PRICE_ANOMALY")
    d = anomaly.details
    assert (d["price_type_as_labelled"], d["workbook_column"], d["old_value"], d["new_value"]) == ("PURCHASE_COST", "Purchase price", "9000", "1290")
    assert d["old_source"] == {"worksheet": "11.09.2026", "row": 7, "date": "2026-09-11"} and d["new_source"]["worksheet"] == "12.09.2026"
    assert d["other_prices_on_that_day"] == {"sales_price": "13000", "purchase_price": "1290"}  # the two price types stay apart
    magic = SourceProduct.objects.get(serial_number=4)
    assert [h["amount"] for h in magic.purchase_prices] == ["9000", "1290"] and [h["amount"] for h in magic.selling_prices] == ["13000"]
    record = HistoricalRecord.objects.get(business_date=SEP12, source_product=magic)
    assert record.purchase_price == 1290 and record.incomplete is True  # not changed to 9,000

    boss = client_for(owner)
    assert decide(boss, anomaly, "DATA_ENTRY_ERROR").data["code"] == "price_required"
    done = decide(boss, anomaly, "DATA_ENTRY_ERROR", note="Typing mistake: 9,000 was meant; confirmed with the supplier invoice", approved_prices=[
        {"source_product": magic.pk, "price_type": "PURCHASE_COST", "amount": "9000", "effective_date": "2026-09-12"}])
    assert done.status_code == 200 and done.data["decision"]["approved_prices"][0]["price_type"] == "PURCHASE_COST"
    record.refresh_from_db()
    assert record.purchase_price == 1290 and record.original["Purchase price"] == "1290"  # the source value is preserved
    assert (record.import_status, record.incomplete) == ("APPROVED", False)
    approved = ApprovedPrice.objects.get()
    assert (approved.price_type, approved.amount, approved.review_item, approved.approved_by) == ("PURCHASE_COST", 9000, anomaly, owner)
    assert AuditLog.objects.filter(action="history.exception_decided", actor=owner).count() == 1

    # the Owner can reopen a decision, with a reason, and the earlier decision stays in the audit trail
    assert boss.post(f"/api/v1/history/review-items/{anomaly.pk}/reopen/", {"reason": "no"}, format="json").status_code == 400
    reopened = boss.post(f"/api/v1/history/review-items/{anomaly.pk}/reopen/", {"reason": "Supplier says the price really did change"}, format="json")
    assert reopened.data["status"] == "PENDING_REVIEW"
    assert AuditLog.objects.get(action="history.exception_reopened").before["resolution"] == "DATA_ENTRY_ERROR"
    # even "the recorded value is correct" needs the price type and applicable amount confirmed
    assert decide(boss, anomaly, "CONFIRMED_CORRECT", note="Promotional cost from the supplier, confirmed").data["code"] == "price_required"
    confirmed = decide(boss, anomaly, "CONFIRMED_CORRECT", note="Promotional cost from the supplier, confirmed", approved_prices=[
        {"source_product": magic.pk, "price_type": "PURCHASE_COST", "amount": "1290", "effective_date": "2026-09-12"}])
    assert confirmed.status_code == 200 and confirmed.data["status"] == "APPROVED"
    assert list(ApprovedPrice.objects.order_by("id").values_list("price_type", "amount")) == [("PURCHASE_COST", 9000), ("PURCHASE_COST", 1290)]  # history kept


# --- 5. mkwawa: no price -----------------------------------------------------------------------

def test_product_without_a_price_is_kept_but_cannot_go_on_sale(owner, cashier, shop, imported):
    missing = item("MISSING_PRICE")
    assert "mkwawa has no sales price" in missing.title and "No price has been invented" in missing.details["statement"]
    mkwawa = SourceProduct.objects.get(normalised_name="mkwawa")
    assert "MISSING_PRICE" in mkwawa.flags and mkwawa.selling_prices == [{"amount": "0", "from": "2026-08-31", "to": "2026-09-12"}]
    assert HistoricalRecord.objects.filter(source_product=mkwawa).count() == 4  # history kept while the mapping is unresolved

    boss = client_for(owner)
    need_identity = boss.post(f"/api/v1/history/source-products/{mkwawa.pk}/decide/", {"approve": True, "create_product": True}, format="json")
    assert need_identity.status_code == 400 and need_identity.data["code"] == "identity_required"
    mapped = boss.post(f"/api/v1/history/source-products/{mkwawa.pk}/decide/", {"approve": True, "create_product": True,
                       "confirmed_name": "Mkwawa water", "confirmed_unit": "Carton", "confirmed_package_size": "12 x 500ml"}, format="json")
    assert mapped.status_code == 200 and mapped.data["review_status"] == "APPROVED" and mapped.data["product_is_active"] is False
    product = Product.objects.get()
    assert (product.name, product.base_unit, product.selling_price, product.is_active) == ("Mkwawa water", "Carton", 0, False)
    assert boss.patch(f"/api/v1/products/{product.pk}/", {"is_active": True}, format="json").status_code == 400  # no activation without a price
    till = client_for(cashier)
    start_shift(till)
    assert sell(till, [(product, 1)], [(shop.cash, 1000)]).data["code"] == "product_inactive"
    assert not StockBalance.objects.filter(product=product).exists()  # no stock came from the Excel either

    assert decide(boss, missing, "PRICE_APPROVED").data["code"] == "price_required"
    price = [{"source_product": mkwawa.pk, "price_type": "SELLING", "amount": "6000", "effective_date": "2026-10-12"}]
    listing = lambda: boss.get("/api/v1/history/source-products/?search=mkwawa").data["results"][0]
    assert listing()["price_unresolved"] is True
    assert decide(boss, missing, "PRICE_APPROVED", approved_prices=price).data["code"] == "identity_required"
    assert decide(boss, missing, "PRICE_APPROVED", approved_prices=price,
                  identities=[{"source_product": mkwawa.pk, "confirmed_name": "Mkwawa water"}]).data["code"] == "identity_required"
    priced = decide(boss, missing, "PRICE_APPROVED", note="Owner confirmed the shelf price for a carton", approved_prices=price,
                    identities=[{"source_product": mkwawa.pk, "confirmed_name": "Mkwawa water", "confirmed_unit": "Carton"}])
    assert priced.status_code == 200 and listing()["price_unresolved"] is False
    assert boss.patch(f"/api/v1/products/{product.pk}/", {"is_active": True, "selling_price": "6000"}, format="json").status_code == 200


def test_owner_maps_a_clean_source_product_and_can_reject_one(owner, accountant, imported):
    boss = client_for(owner)
    safari, konyagi = SourceProduct.objects.get(serial_number=6), SourceProduct.objects.get(serial_number=1)
    assert safari.flags == ["UNIT_NOT_STATED"] and safari.proposed_unit == ""  # the unit is never guessed
    body = {"approve": True, "create_product": True, "confirmed_name": "Safari Kubwa", "confirmed_unit": "Crate", "selling_price": "36000", "cost_price": "34000"}
    assert client_for(accountant).post(f"/api/v1/history/source-products/{safari.pk}/decide/", body, format="json").status_code == 403
    done = boss.post(f"/api/v1/history/source-products/{safari.pk}/decide/", body, format="json")
    assert done.status_code == 200 and done.data["product_name"] == "Safari Kubwa" and done.data["product_is_active"] is True
    product = Product.objects.get()
    assert (product.selling_price, product.cost_price, product.base_unit) == (36000, 34000, "Crate")
    assert {(p.price_type, p.amount) for p in ApprovedPrice.objects.filter(source_product=safari)} == {("SELLING", 36000), ("PURCHASE_COST", 34000)}
    assert stock_of(product) == 0  # mapping a product never gives it stock
    assert boss.post(f"/api/v1/history/source-products/{konyagi.pk}/decide/", {"approve": False}, format="json").data["code"] == "note_required"
    rejected = boss.post(f"/api/v1/history/source-products/{konyagi.pk}/decide/", {"approve": False, "note": "No longer stocked"}, format="json")
    assert rejected.data["review_status"] == "REJECTED" and HistoricalRecord.objects.filter(source_product=konyagi).count() == 4
    assert AuditLog.objects.filter(action="history.mapping_decided").count() == 2


# --- 6. Unlabelled totals rows -------------------------------------------------------------------

def test_unlabelled_rows_are_preserved_and_only_ever_labelled(owner, accountant, imported):
    rows = list(UnclassifiedRow.objects.filter(business_date=AUG31).order_by("source_row"))
    assert [(r.worksheet, r.source_row, r.proposed_classification) for r in rows] == [
        ("31.08.2026", 11, "Column totals"), ("31.08.2026", 12, "Not known")]
    assert rows[0].original == [{"column": "J", "value": "45000"}, {"column": "K", "value": "7500"}]  # exactly as written
    assert "Total Sales" in rows[0].proposal_basis and rows[1].original == [{"column": "J", "value": "632700"}]
    assert all(r.status == "PENDING_REVIEW" and r.classification == "" for r in UnclassifiedRow.objects.all())  # proposal not applied

    office = client_for(accountant)
    listed = office.get("/api/v1/history/unclassified-rows/?row=12").data
    assert listed["count"] == 4 and listed["results"][0]["proposal_basis"]
    totals = list(UnclassifiedRow.objects.filter(source_row=11).values_list("pk", flat=True))
    unknown = list(UnclassifiedRow.objects.filter(source_row=12).values_list("pk", flat=True))
    url = "/api/v1/history/unclassified-rows/decide/"
    assert office.post(url, {"ids": totals, "action": "CLASSIFY"}, format="json").data["code"] == "classification_required"
    assert office.post(url, {"ids": totals, "action": "CLASSIFY", "classification": "Column totals"}, format="json").data == {"updated": 4}
    assert office.post(url, {"ids": unknown, "action": "EXCLUDE"}, format="json").data["code"] == "reason_required"
    assert office.post(url, {"ids": unknown[:2], "action": "EXCLUDE", "reason": "Owner's private note, not shop data"}, format="json").status_code == 200
    assert office.post(url, {"ids": unknown[2:], "action": "FLAG"}, format="json").status_code == 200
    # changing an earlier decision is a correction and needs a reason
    assert office.post(url, {"ids": totals[:1], "action": "FLAG"}, format="json").data["code"] == "reason_required"
    assert office.post(url, {"ids": totals[:1], "action": "FLAG", "reason": "Does not match on closer look"}, format="json").status_code == 200
    statuses = dict(UnclassifiedRow.objects.values_list("pk", "status"))
    assert statuses[totals[1]] == "APPROVED" and statuses[unknown[0]] == "REJECTED" and statuses[unknown[3]] == "FLAGGED" and statuses[totals[0]] == "FLAGGED"
    assert AuditLog.objects.filter(action="history.row_decided").count() == 9
    shown = next(r for r in office.get("/api/v1/history/unclassified-rows/?row=11").data["results"] if r["id"] == totals[0])
    assert (shown["worksheet"], shown["source_row"], shown["decided_by_name"]) == ("31.08.2026", 11, "Accountant")
    assert [(h["status_label"], h["classification"], h["reason"], h["reviewer"]) for h in shown["history"]] == [
        ("Classified", "Column totals", "", "Accountant"), ("Needs investigation", "Column totals", "Does not match on closer look", "Accountant")]
    assert all(h["decided_at"] for h in shown["history"])
    # whatever the label, the rows never became products, records or transactions
    assert HistoricalRecord.objects.count() == 24 and SourceProduct.objects.count() == 6 and not Sale.objects.exists()
    assert UnclassifiedRow.objects.get(pk=unknown[0]).original == [{"column": "J", "value": "632700"}]


# --- Reconciliation status, dashboard, exception report -----------------------------------------

def resolve_everything(owner):
    boss = client_for(owner)
    decide(boss, item("MISSING_DATES"), "ACCEPTED_INCOMPLETE")
    henkens = list(SourceProduct.objects.filter(normalised_name="henken"))
    magic = SourceProduct.objects.get(serial_number=4)
    answers = [
        decide(boss, item("SAME_NAME_DIFFERENT_PRICE"), "SAME_PRODUCT_PRICE_CHANGE",
               identities=[{"source_product": h.pk, "confirmed_name": "Heineken", "confirmed_package_size": "330ml"} for h in henkens],
               approved_prices=[{"source_product": h.pk, "price_type": "SELLING", "amount": "4000", "effective_date": "2026-08-31"} for h in henkens]),
        decide(boss, item("PRICE_ANOMALY"), "CONFIRMED_CORRECT", approved_prices=[
            {"source_product": magic.pk, "price_type": "PURCHASE_COST", "amount": "1290", "effective_date": "2026-09-12"}]),
        decide(boss, item("MISSING_PRICE"), "NOT_FOR_SALE"),
    ]
    assert [a.status_code for a in answers] == [200, 200, 200], [a.data for a in answers]
    for sp in SourceProduct.objects.all():
        boss.post(f"/api/v1/history/source-products/{sp.pk}/decide/", {"approve": False, "note": "History only"}, format="json")
    boss.post("/api/v1/history/unclassified-rows/decide/", {"ids": list(UnclassifiedRow.objects.values_list("pk", flat=True)),
              "action": "EXCLUDE", "reason": "Not transactional data"}, format="json")
    return boss


def test_history_is_not_reconciled_until_everything_is_decided_and_the_owner_accepts(owner, accountant, imported):
    boss = client_for(owner)
    board = boss.get("/api/v1/history/dashboard/").data
    assert board["reconciled"] is False and board["statement"].startswith("NOT RECONCILED")
    assert board["blockers"] == {"exceptions": 4, "unlabelled_rows": 8, "product_mappings": 6}
    assert board["missing_dates"][0] == "2026-09-01" and len(board["missing_dates"]) == 9
    assert board["records"]["total"] == 24 and board["records"]["incomplete"] > 0 and board["can_be_accepted"] is False
    early = boss.post("/api/v1/history/accept/", {"note": "Looks fine to me overall"}, format="json")
    assert early.status_code == 409 and early.data["code"] == "still_pending" and early.data["details"]["exceptions"] == 4

    resolve_everything(owner)
    board = boss.get("/api/v1/history/dashboard/").data
    assert not any(board["blockers"].values()) and board["can_be_accepted"] is True
    assert board["reconciled"] is False  # deciding everything is not the same as accepting it
    assert client_for(accountant).post("/api/v1/history/accept/", {"note": "Accepting on the owner's behalf"}, format="json").status_code == 403
    assert boss.post("/api/v1/history/accept/", {"note": "ok"}, format="json").status_code == 400
    accepted = boss.post("/api/v1/history/accept/", {"note": "All exceptions reviewed with the accountant on 12 Oct"}, format="json")
    assert accepted.status_code == 200 and accepted.data["reconciled"] is True and accepted.data["acceptance"]["by"] == "Owner"
    assert HistoryAcceptance.objects.get().snapshot["records"]["total"] == 24
    # figures accepted without explanation stay marked incomplete even after acceptance
    assert accepted.data["records"]["incomplete"] >= 1

    # anything new afterwards withdraws the claim until the Owner accepts again
    boss.post(f"/api/v1/history/review-items/{item('PRICE_ANOMALY').pk}/reopen/", {"reason": "Looking at this again"}, format="json")
    board = boss.get("/api/v1/history/dashboard/").data
    assert board["reconciled"] is False and board["acceptance"]["still_valid"] is False


def test_exception_report_download(owner, accountant, cashier, imported):
    assert client_for(cashier).get("/api/v1/history/exception-report/").status_code == 403
    response = client_for(accountant).get("/api/v1/history/exception-report/")
    assert response.status_code == 200 and response["Content-Type"].startswith("text/csv") and "attachment" in response["Content-Disposition"]
    text = response.content.decode("utf-8-sig")
    assert "NOT RECONCILED" in text and "Same name, different prices" in text and "Missing: 2026-09-01" in text
    assert "Purchase price: 9000 (11.09.2026 row 7) -> 1290 (12.09.2026 row 7)" in text
    assert "S/N 5: mkwawa" in text and "31.08.2026 row 12" in text and "J=632700" in text
    records = client_for(accountant).get("/api/v1/history/records/?incomplete=true")
    assert records.status_code == 200 and records.data["count"] > 0 and records.data["results"][0]["worksheet"]


# --- 1. Supplementary workbook --------------------------------------------------------------------

def supplementary() -> bytes:
    """The missing 1-9 September sheets, plus 10 Sep again with a different figure."""
    def day(konyagi_open, konyagi_sold):
        return [line(1, "Konyagi  ndogo", konyagi_open, konyagi_sold, 3750, 4500), line(2, "Henken", 102, 0, 2917, 3500),
                line(3, "Henken", 0, 0, 3125, 4000), line(4, "MAGIC MOMENT kati", 48, 0, 9000, 13000), line(5, "mkwawa", 0, 0, 0, 0),
                line(6, "Safari Kubwa", 20, 0, 34000, 36000)]
    days = {date(2026, 9, 1): day(90, 20)}
    days.update({date(2026, 9, d): day(70, 0) for d in range(2, 10)})
    days[SEP10] = day(70, 6)  # the stored record says 5 were sold
    return workbook(days, with_totals=False)


def test_supplementary_workbook_adds_only_what_is_new(owner, imported):
    boss = client_for(owner)
    preview = upload(boss, supplementary(), "/api/v1/history/imports/dry-run/", name="September.xlsx").data
    assert (preview["summary"]["accepted_new"], preview["summary"]["duplicates_skipped"], preview["summary"]["conflicts_not_imported"]) == (54, 5, 1)
    assert "MISSING_DATES" not in {e["kind"] for e in preview["exceptions"]}  # the preview already shows the gap would close
    assert HistoricalRecord.objects.count() == 24

    done = commit(boss, supplementary(), name="September.xlsx")
    assert done.status_code == 201 and done.data["batch"]["is_supplementary"] is True
    assert HistoricalRecord.objects.count() == 78 and SourceProduct.objects.count() == 6  # no duplicates
    assert HistoricalRecord.objects.filter(business_date=SEP10).count() == 6
    gap = item("MISSING_DATES")
    assert (gap.status, gap.resolution) == ("APPROVED", "NO_LONGER_APPLIES") and "September.xlsx" in gap.note
    konyagi = HistoricalRecord.objects.get(business_date=SEP10, source_product__serial_number=1)
    assert konyagi.sold_qty == 5 and konyagi.source_file == "MERIBA STORE.xlsx"  # the stored record was not overwritten
    clash = item("DUPLICATE_CONFLICT")
    line_ = clash.details["lines"][0]
    assert (line_["already_stored"]["sold_qty"], line_["in_this_file"]["sold_qty"], line_["in_this_file"]["worksheet"]) == ("5", "6", "10.09.2026")
    assert konyagi.import_status == "PENDING_REVIEW" and konyagi.incomplete is True
    assert commit(boss, supplementary(), name="September.xlsx").data["nothing_new"] is True


def test_a_sheet_that_is_not_a_day_is_reported_not_guessed_at(owner):
    data = workbook({AUG31: [line(1, "Konyagi", 10, 1, 3750, 4500)]}, extra_sheets=["Summary"])
    response = commit(client_for(owner), data)
    assert response.status_code == 201 and response.data["batch"]["summary"]["worksheets_not_read"] == 1
    unread = item("UNREADABLE_SHEET")
    assert "Summary" in unread.title and "not a date" in unread.details["reason"]
    assert Branch.objects.count() == 1 and HistoricalRecord.objects.count() == 1


def test_lines_that_do_not_add_up_or_cannot_be_read_are_flagged_as_written(owner):
    bad = line(1, "Konyagi", 10, 2, 3750, 4500)
    bad[6] = 9  # closing stock should be 8
    unreadable = line(2, "Safari", 5, 0, 34000, 36000)
    unreadable[8] = "ask boss"
    response = commit(client_for(owner), workbook({AUG31: [bad, unreadable]}))
    assert response.data["batch"]["summary"]["rows_with_warnings"] == 1 and response.data["batch"]["summary"]["rows_with_errors"] == 1
    konyagi, safari = HistoricalRecord.objects.order_by("source_row")
    assert konyagi.validation_status == "WARNING" and "closing stock says 9" in konyagi.validation_messages[0] and konyagi.closing_qty == 9
    assert safari.validation_status == "ERROR" and safari.sales_price is None and safari.original["Sales price"] == "ask boss"
    assert {i.kind for i in ReviewItem.objects.all()} >= {"ARITHMETIC", "INVALID_VALUES"}


# --- 7. Opening stock comes from a physical count, never from the Excel ---------------------------

def test_excel_closing_stock_is_only_a_reference_on_the_count_sheet(owner, accountant, shop, imported):
    shop.cfg.went_live_at = None
    shop.cfg.opening_stock_confirmed_at = None
    shop.cfg.save()
    boss, prep = client_for(owner), client_for(accountant)
    safari = SourceProduct.objects.get(serial_number=6)
    boss.post(f"/api/v1/history/source-products/{safari.pk}/decide/", {"approve": True, "create_product": True, "confirmed_name": "Safari Kubwa",
              "confirmed_unit": "Crate", "selling_price": "36000"}, format="json")
    product = Product.objects.get()
    assert stock_of(product) == 0 and not InventoryMovement.objects.exists()  # 20 crates in the Excel; none in live stock

    sheet = prep.post("/api/v1/stock/opening-counts/", {"request_key": str(uuid.uuid4()), "count_date": "2026-10-31",
                      "counted_by_name": "Asha", "lines": [{"product": product.pk, "quantity": "17"}]}, format="json")
    counted = sheet.data["lines"][0]
    assert (counted["unit"], counted["counted_quantity"], counted["reference_quantity"], counted["variance"]) == ("Crate", "17.000", "20.000", "-3.000")
    assert counted["reference_source"] == "Excel closing stock on 12 Sep 2026 (reference only)"
    assert stock_of(product) == 0  # still nothing live: counted but not yet approved
    approved = boss.post(f"/api/v1/stock/opening-counts/{sheet.data['id']}/review/", {"approve": True}, format="json").data
    assert approved["lines"][0]["approved_quantity"] == "17.000" and approved["lines"][0]["movement"]
    assert stock_of(product) == 17  # the physical count, not the Excel figure of 20
    assert InventoryMovement.objects.get().source_type == "OpeningCount"


# --- The real workbook, when it is on this computer ------------------------------------------------

REAL = os.environ.get("MERIBA_WORKBOOK", r"C:\Users\KAROLI\OneDrive\Desktop\MERIBA-Store-System\MERIBA  STORE OCT  2026.xlsx")


@pytest.mark.skipif(not os.path.exists(REAL), reason="the shop's workbook is not on this computer")
def test_real_workbook_dry_run_then_import(owner, shop):
    try:
        data = open(REAL, "rb").read()
    except PermissionError:
        pytest.skip("the workbook is open in Excel, which locks it; close it to run this test")
    boss = client_for(owner)
    preview = upload(boss, data, "/api/v1/history/imports/dry-run/", name="MERIBA STORE OCT 2026.xlsx").data
    s = preview["summary"]
    assert (s["worksheets_read"], s["product_rows"], s["date_from"], s["date_to"]) == (31, 3503, "2026-08-31", "2026-10-09")
    assert s["workbook_total_sales"] == "50395300" and s["rows_with_errors"] == 0
    assert not HistoricalRecord.objects.exists()

    assert commit(boss, data, name="MERIBA STORE OCT 2026.xlsx").status_code == 201
    assert HistoricalRecord.objects.count() == 3503 and UnclassifiedRow.objects.count() == 434
    assert not InventoryMovement.objects.exists() and not StockBalance.objects.exists() and not Product.objects.exists()
    gap = ReviewItem.objects.get(kind="MISSING_DATES")
    assert gap.details["missing_dates"][0] == "2026-09-01" and len(gap.details["missing_dates"]) == 9 and len(gap.details["affected_products"]) == 57
    henken = ReviewItem.objects.get(key="SAME_NAME|henken")
    assert [(l["serial"], l["sales_price_history"][0]["amount"]) for l in henken.details["lines"]] == [(4, "3500"), (5, "4000")]
    magic = ReviewItem.objects.get(kind="PRICE_ANOMALY")
    assert (magic.details["old_value"], magic.details["new_value"], magic.details["price_type_as_labelled"]) == ("9000", "1290", "PURCHASE_COST")
    assert "mkwawa" in ReviewItem.objects.get(kind="MISSING_PRICE").title.lower()
    assert set(UnclassifiedRow.objects.values_list("source_row", flat=True)) == set(range(117, 124)) | set(range(125, 132))
    assert commit(boss, data, name="MERIBA STORE OCT 2026.xlsx").data["nothing_new"] is True
    assert HistoricalRecord.objects.count() == 3503
    print("\nReal workbook exceptions:", sorted(ReviewItem.objects.values_list("title", flat=True)))
    print("Unlabelled row proposals by row:", sorted(set(UnclassifiedRow.objects.values_list("source_row", "proposed_classification"))))


# --- Controls confirmed as policy on 12 Oct 2026 ---------------------------------------------------

def test_exception_report_is_available_from_the_dry_run_before_anything_is_committed(owner, accountant, cashier):
    url = "/api/v1/history/imports/dry-run/?download=csv"
    assert upload(client_for(cashier), primary(), url).status_code == 403
    response = upload(client_for(accountant), primary(), url)
    assert response.status_code == 200 and response["Content-Type"].startswith("text/csv") and "attachment" in response["Content-Disposition"]
    text = response.content.decode("utf-8-sig")
    assert text.startswith("MERIBA historical import PREVIEW - nothing has been saved") and "NOT RECONCILED" in text
    assert "Same name, different prices" in text and "Missing: 2026-09-01" in text and "Unusual price change" in text
    assert "mkwawa has no sales price" in text and "Row 12 in 4 sheet(s)" in text and "Never imported as a transaction" in text
    for model in (ImportBatch, HistoricalRecord, SourceProduct, ReviewItem, UnclassifiedRow):
        assert not model.objects.exists(), model.__name__


def test_go_live_waits_for_price_questions_but_not_for_the_rest_of_the_history_review(owner, shop, make_user, make_product, imported):
    """Development and review of the old records can go on in parallel; only identity and price questions hold up go-live."""
    shop.cfg.went_live_at = None
    shop.cfg.opening_stock_confirmed_at = None
    shop.cfg.save()
    boss = client_for(owner)
    make_user("till1", "CASHIER")
    product = make_product("Konyagi ndogo", "4500")
    sheet = boss.post("/api/v1/stock/opening-counts/", {"request_key": str(uuid.uuid4()), "count_date": "2026-10-31", "counted_by_name": "Asha",
                      "lines": [{"product": product.pk, "quantity": "40"}]}, format="json")
    boss.post(f"/api/v1/stock/opening-counts/{sheet.data['id']}/review/", {"approve": True}, format="json")
    go = {"opening_cumulative_sales": "0", "as_of_date": "2026-10-31", "confirm": True, "confirm_new_baseline": True}

    step = lambda: next(s for s in boss.get("/api/v1/setup/status/").data["steps"] if s["key"] == "price_questions")
    assert step()["done"] is False and "3 question(s) waiting for the Owner" in step()["detail"]
    blocked = boss.post("/api/v1/setup/go-live/", go, format="json")
    assert blocked.status_code == 409 and blocked.data["code"] == "setup_incomplete"
    assert any("price questions" in title for title in blocked.data["details"])

    henkens = list(SourceProduct.objects.filter(normalised_name="henken"))
    magic = SourceProduct.objects.get(serial_number=4)
    decide(boss, item("SAME_NAME_DIFFERENT_PRICE"), "DIFFERENT_VARIANTS",
           identities=[{"source_product": h.pk, "confirmed_name": "Heineken", "confirmed_package_size": s} for h, s in zip(henkens, ("330ml", "500ml"))],
           approved_prices=[{"source_product": h.pk, "price_type": "SELLING", "amount": a, "effective_date": "2026-08-31"} for h, a in zip(henkens, ("3500", "4000"))])
    decide(boss, item("PRICE_ANOMALY"), "CONFIRMED_CORRECT", approved_prices=[
        {"source_product": magic.pk, "price_type": "PURCHASE_COST", "amount": "1290", "effective_date": "2026-09-12"}])
    decide(boss, item("MISSING_PRICE"), "NOT_FOR_SALE")
    assert step()["done"] is True

    # missing dates, product mappings and unlabelled rows are still open, and the history is still not reconciled
    board = boss.get("/api/v1/history/dashboard/").data
    assert board["reconciled"] is False and board["blockers"]["exceptions"] == 1 and board["blockers"]["unlabelled_rows"] == 8
    live = boss.post("/api/v1/setup/go-live/", go, format="json")
    assert live.status_code == 200 and live.data["is_live"] is True
    assert boss.get("/api/v1/history/dashboard/").data["statement"].startswith("NOT RECONCILED")  # going live claims nothing about the past
    assert stock_of(product) == 40 and InventoryMovement.objects.count() == 1


def test_statuses_are_distinct_and_named_as_agreed(owner, imported):
    from history.models import ReviewStatus

    assert [label for _, label in ReviewStatus.choices] == ["Pending review", "Approved", "Rejected", "Imported"]
    boss = client_for(owner)
    decide(boss, item("MISSING_DATES"), "REJECTED_RECORDS")
    statuses = set(HistoricalRecord.objects.values_list("import_status", flat=True))
    assert statuses == {"IMPORTED", "PENDING_REVIEW", "REJECTED"}  # clean lines, lines under open exceptions, rejected lines
    labels = {r["import_status_label"] for r in boss.get("/api/v1/history/records/?page_size=100").data["results"]}
    assert labels == {"Imported", "Pending review", "Rejected"}
    assert ReviewItem.objects.get(kind="MISSING_DATES").status == "REJECTED"