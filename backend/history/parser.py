"""Reads the shop's daily-sheet Excel workbook. Read-only: it only describes what is in the file.

Layout it understands (one worksheet per day, named DD.MM.YYYY):
    S/N | Details | [pack] | [Description] | <opening> | Add Purchases | Total stock | Sales | Closing stock |
    Purchase price | Sales price | Total Sales | Profit | Stock | T.Purchases
The opening-stock column has no proper heading; it is the column just before "Add Purchases".
Rows under the product list that have no S/N and name are kept as "other rows" and never interpreted.
"""
import hashlib
import zipfile
from datetime import datetime
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.utils import get_column_letter

from core.errors import BusinessRuleError

MAX_ARCHIVE_MEMBERS = 1000
MAX_EXPANDED_BYTES = 100 * 1024 * 1024

# heading text (lower-case, single spaces) -> our field name. Price types are kept apart on purpose.
HEADINGS = {
    "add purchases": "purchases_qty", "total stock": "total_stock", "sales": "sold_qty", "closing stock": "closing_qty",
    "purchase price": "purchase_price", "sales price": "sales_price", "total sales": "total_sales", "profit": "profit",
    "stock": "stock_value", "t.purchases": "total_purchases",
}
NUMERIC_FIELDS = ["opening_qty", *HEADINGS.values()]
REQUIRED = ["purchases_qty", "sold_qty", "closing_qty", "purchase_price", "sales_price"]


def squash(value) -> str:
    return " ".join(str(value).split())


def normalise_name(value) -> str:
    return squash(value).lower()


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return squash(value)


def _number(value):
    """Returns (Decimal or None, error or None). An empty cell is None, not zero."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    if isinstance(value, bool):
        return None, "not a number"
    try:
        number = Decimal(str(value).replace(",", "").strip())
    except InvalidOperation:
        return None, f"'{squash(value)}' is not a number"
    if not number.is_finite():
        return None, "not a number"
    return number, None


def file_digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def parse_workbook(stream, filename: str) -> dict:
    try:
        with zipfile.ZipFile(stream) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS or sum(member.file_size for member in members) > MAX_EXPANDED_BYTES:
                raise BusinessRuleError("The workbook expands beyond the supported size.", code="too_large", status_code=400)
    finally:
        stream.seek(0)
    workbook = openpyxl.load_workbook(stream, read_only=True, data_only=True)
    sheets = []
    for name in workbook.sheetnames:
        ws = workbook[name]
        sheet = {"name": name, "date": None, "error": None, "headings": {}, "products": [], "other_rows": []}
        sheets.append(sheet)
        try:
            sheet["date"] = datetime.strptime(name.strip(), "%d.%m.%Y").date()
        except ValueError:
            # FR-BR-03: a sheet name is only ever a date. Anything else is reported, never guessed at.
            sheet["error"] = "The sheet name is not a date in the form DD.MM.YYYY, so the sheet was not read."
            continue
        rows = list(ws.iter_rows(values_only=True))
        header_index = next((i for i, r in enumerate(rows) if r and r[0] is not None and squash(r[0]).upper() == "S/N"), None)
        if header_index is None:
            sheet["error"] = "No heading row starting with 'S/N' was found, so the sheet was not read."
            continue
        columns = {}
        for index, cell in enumerate(rows[header_index]):
            field = HEADINGS.get(normalise_name(cell)) if cell is not None else None
            if field and field not in columns:
                columns[field] = index
                sheet["headings"][field] = squash(cell)
        missing = [f for f in REQUIRED if f not in columns]
        if missing:
            sheet["error"] = "Expected column heading(s) not found: " + ", ".join(missing) + "."
            continue
        columns["opening_qty"] = columns["purchases_qty"] - 1
        sheet["headings"]["opening_qty"] = _text(rows[header_index][columns["opening_qty"]]) or "(no heading)"
        extra_text_columns = list(range(2, columns["opening_qty"]))  # pack / description, when present

        for offset, row in enumerate(rows[header_index + 1:]):
            row_number = header_index + 2 + offset
            if not any(c is not None and str(c).strip() for c in row):
                continue
            serial, serial_error = _number(row[0])
            name_cell = _text(row[1]) if len(row) > 1 else ""
            if serial is None or serial_error or serial != serial.to_integral_value() or not name_cell:
                sheet["other_rows"].append({
                    "row": row_number,
                    "cells": [{"column": get_column_letter(i + 1), "value": _text(c)} for i, c in enumerate(row)
                              if c is not None and str(c).strip()],
                })
                continue
            product = {
                "row": row_number, "serial": int(serial), "name": name_cell, "normalised": normalise_name(name_cell),
                "pack": _text(row[extra_text_columns[0]]) if extra_text_columns else "",
                "description": _text(row[extra_text_columns[1]]) if len(extra_text_columns) > 1 else "",
                "original": {"S/N": _text(row[0]), "Details": name_cell}, "errors": [],
            }
            for index in extra_text_columns:
                product["original"][f"Column {get_column_letter(index + 1)}"] = _text(row[index])
            for field in NUMERIC_FIELDS:
                index = columns.get(field)
                cell = row[index] if index is not None and index < len(row) else None
                product["original"][sheet["headings"].get(field, field) if field != "opening_qty" else "Opening (column before Add Purchases)"] = _text(cell)
                value, error = _number(cell)
                product[field] = value
                if error:
                    product["errors"].append(f"{sheet['headings'].get(field, field)}: {error}")
            sheet["products"].append(product)
    workbook.close()
    return {"filename": filename, "sheets": sheets}
